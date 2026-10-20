import os,sys,re,html,subprocess,socket,math,json,ipaddress
import tkinter as tk
import tkinter.simpledialog
from tkinter import ttk,filedialog,messagebox
from pathlib import Path
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.barcode import eanbc,code128
from reportlab.graphics.shapes import Drawing
from reportlab.graphics import renderPDF

BASE=Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parent))
SETTINGS_DIR=(Path.home()/'Library'/'Application Support'/'TheVandalStickerprogramma') if sys.platform=='darwin' else Path(os.environ.get('APPDATA',str(Path.home()))) / 'TheVandalStickerprogramma'
SETTINGS_FILE=SETTINGS_DIR / 'printerinstellingen.json'
DEFAULT_SETTINGS={'ip':'10.60.1.231','port':'9100','gap':'3','pitch':'32'}
LAST_EXCEL_FILE=SETTINGS_DIR / 'laatste_excel.json'
def get_last_excel():
    try:
        path=json.loads(LAST_EXCEL_FILE.read_text(encoding='utf-8')).get('path','')
        return path if path and Path(path).is_file() else None
    except (OSError,ValueError,TypeError,AttributeError):
        return None
def remember_excel(path):
    try:
        SETTINGS_DIR.mkdir(parents=True,exist_ok=True)
        LAST_EXCEL_FILE.write_text(json.dumps({'path':str(Path(path).resolve())}),encoding='utf-8')
    except OSError:
        pass
def load_settings():
    try:
        data=json.loads(SETTINGS_FILE.read_text(encoding='utf-8'))
        if isinstance(data,dict):return {k:str(data.get(k,v)) for k,v in DEFAULT_SETTINGS.items()}
    except (OSError,ValueError,TypeError):pass
    return DEFAULT_SETTINGS.copy()

FONTDIR=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'
def init_fonts():
    paths=[(FONTDIR/'arial.ttf',FONTDIR/'arialbd.ttf'),(Path('/Library/Fonts/Arial.ttf'),Path('/Library/Fonts/Arial Bold.ttf')),(Path('/System/Library/Fonts/Supplemental/Arial.ttf'),Path('/System/Library/Fonts/Supplemental/Arial Bold.ttf')),(Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))]
    for normal,bold in paths:
        if normal.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont('LabelRegular',str(normal)))
            pdfmetrics.registerFont(TTFont('LabelBold',str(bold)))
            return
    raise RuntimeError('Geen bruikbaar TrueType-lettertype gevonden')

def value(x):
    if x is None:return ''
    if isinstance(x,float) and x.is_integer():return str(int(x))
    return str(x).strip()

def valid_ean13(s):
    if not re.fullmatch(r'\d{13}',s):return False
    return (10-sum(int(v)*(1 if i%2==0 else 3) for i,v in enumerate(s[:12]))%10)%10==int(s[-1])

def load_products(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    sh=wb.active
    rows=sh.iter_rows(values_only=True)
    headers=[value(x).lower() for x in next(rows)]
    aliases={'sku':['code','sku','artikelnummer'],'description':['omschrijving','description','name','naam'],'size':['size','maat'],'barcode':['barcode','ean','ean13'],'color':['color','kleur']}
    indexes={k:next((headers.index(a) for a in opts if a in headers),None) for k,opts in aliases.items()}
    if indexes['sku'] is None or indexes['description'] is None:raise ValueError('Excel moet kolommen Code en Omschrijving bevatten.')
    products=[]
    for r in rows:
        d={k:value(r[i]) if i is not None and i<len(r) else '' for k,i in indexes.items()}
        if d['sku'] or d['description']:
            d['barcode']=re.sub(r'\.0$','',d['barcode'])
            products.append(d)
    wb.close()
    return products

def fitted(text,font,width,initial,minimum):
    size=initial
    while size>minimum and pdfmetrics.stringWidth(text,font,size)>width:size-=.25
    if pdfmetrics.stringWidth(text,font,size)>width:
        while text and pdfmetrics.stringWidth(text+'…',font,size)>width:text=text[:-1]
        text+='…'
    return text,size

LABEL_W, LABEL_H = 45, 30
GAP_MM = 3.0  # horizontale tussenruimte; afstellen met testafdruk
PITCH_MM = 32.0  # verticale rijafstand: 30 mm etiket + 2 mm tussenruimte

def draw_label(c,p):
    # PDF preview uses the same reserved vertical zones as the ZPL layout.
    W,H=LABEL_W*mm,LABEL_H*mm
    desc=html.unescape(p['description'])
    words=desc.split()
    lines=[]
    for word in words:
        if not lines or pdfmetrics.stringWidth(lines[-1]+' '+word,'LabelBold',8.1) > W-4*mm:
            lines.append(word)
        else:lines[-1]+=' '+word
    if len(lines)>2:
        lines=lines[:2]
        lines[1],_=fitted(lines[1]+'…','LabelBold',W-4*mm,8.1,7.0)
    for n,line in enumerate(lines):
        line,fs=fitted(line,'LabelBold',W-4*mm,8.1,6.2)
        c.setFont('LabelBold',fs)
        c.drawCentredString(W/2,H-(3.5+n*3.2)*mm,line)
    bc=p['barcode']
    if valid_ean13(bc):
        bar=eanbc.Ean13BarcodeWidget(bc,barHeight=8.5*mm,barWidth=.26*mm,humanReadable=True)
        bounds=bar.getBounds();bw=bounds[2]-bounds[0]
        d=Drawing(bw,11*mm);d.add(bar)
        renderPDF.draw(d,c,(W-bw)/2,10.6*mm)
    else:
        bar=code128.Code128(bc,barHeight=8.5*mm,barWidth=.20*mm,humanReadable=False)
        bw=bar.width
        c.saveState();c.translate((W-min(bw,W-4*mm))/2,12.0*mm)
        if bw>W-4*mm:c.scale((W-4*mm)/bw,1)
        bar.drawOn(c,0,0);c.restoreState()
        digits,ds=fitted(bc,'LabelRegular',W-4*mm,6.6,5)
        c.setFont('LabelRegular',ds);c.drawCentredString(W/2,10.0*mm,digits)
    sku,sz=fitted('SKU: '+p['sku'],'LabelBold',W-3*mm,8.4,6.2)
    c.setFont('LabelBold',sz);c.drawCentredString(W/2,6.2*mm,sku)
    left,ls=fitted('Size: '+(p['size'] or '-'),'LabelBold',14*mm,8.0,6.3)
    right,rs=fitted('Color: '+(p['color'] or '-'),'LabelBold',27*mm,8.0,6.0)
    c.setFont('LabelBold',ls);c.drawString(1.5*mm,1.6*mm,left)
    c.setFont('LabelBold',rs);c.drawString(16*mm,1.6*mm,right)


def expanded(items):
    for product,quantity in items:
        for _ in range(quantity):
            yield product

def make_pdf(items,path):
    init_fonts()
    page_w=(LABEL_W*2+GAP_MM)*mm
    page_h=PITCH_MM*mm
    c=canvas.Canvas(path,pagesize=(page_w,page_h),pageCompression=1)
    labels=list(expanded(items))
    for i in range(0,len(labels),2):
        for column,p in enumerate(labels[i:i+2]):
            c.saveState()
            c.translate((LABEL_W+GAP_MM)*mm*column,(PITCH_MM-LABEL_H)*mm)
            draw_label(c,p)
            c.restoreState()
        c.showPage()
    c.save()

def zpl_row(left,right=None,dpi=203,gap_mm=GAP_MM,pitch_mm=PITCH_MM):
    """One ZPL format per physical row, two 45 x 30 mm labels."""
    dot=lambda x:round(x*dpi/25.4)
    width=dot(2*LABEL_W+gap_mm)
    height=dot(pitch_mm)
    def field(p,offset):
        if p is None:return ''
        x=lambda mm:dot(offset+mm)
        y=dot
        desc=zpl_safe(html.unescape(p['description']))
        sku=zpl_safe('SKU: '+p['sku'])
        size=zpl_safe('Size: '+(p['size'] or '-'))
        color=zpl_safe('Color: '+(p['color'] or '-'))
        bc=zpl_safe(p['barcode'])
        # Native Zebra EAN-13: two-dot modules at 203 dpi, valid checksum,
        # and a quiet zone on both sides.  Avoid 1-dot Code128 bars.
        if valid_ean13(bc):
            # EAN-13 encodes 12 data digits; Zebra calculates the check digit.
            # 95 modules * 2 dots = 190 dots (~23.8 mm) + quiet zones.
            barcode_cmd=(f'^FO{x(8.5)},{y(9.3)}^BY2,2^BEN,{y(9.2)},Y,N'
                         f'^FD{bc[:12]}^FS')
        else:
            # Other barcode formats: use Code128 at >=2 dots per module.
            # Refuse codes that cannot fit with 3 mm quiet zones.
            from reportlab.graphics.barcode.code128 import Code128
            enc=Code128(bc,barWidth=1,barHeight=10,quiet=False)
            modules=round(enc.width)
            available=dot(39)
            if modules*2>available:
                raise ValueError(f'Barcode {bc} is te lang voor een scanbare 45 mm sticker. '
                                 'Gebruik een geldig EAN-13 nummer of een groter etiket.')
            start_mm=(45-modules*2*25.4/dpi)/2
            barcode_cmd=(f'^FO{x(start_mm)},{y(9.3)}^BY2,2^BCN,{y(9.2)},Y,N,N^FD{bc}^FS')
        return (f'^FO{x(1.5)},{y(1.0)}^A0N,{y(3.1)},{y(2.5)}^FB{dot(42)},2,{y(0.15)},C,0^FD{desc}^FS'
                +barcode_cmd+
                f'^FO{x(1.2)},{y(22.0)}^A0N,{y(3.0)},{y(2.2)}^FB{dot(42.5)},1,0,C,0^FD{sku}^FS'
                f'^FO{x(1.5)},{y(26.0)}^A0N,{y(2.9)},{y(2.1)}^FB{dot(14)},1,0,L,0^FD{size}^FS'
                f'^FO{x(16)},{y(26.0)}^A0N,{y(2.9)},{y(2.0)}^FB{dot(27)},1,0,L,0^FD{color}^FS')
    return (f'^XA^CI28^PW{width}^LL{height}^LH0,0^LS0^LT0'
            +field(left,0)+field(right,LABEL_W+gap_mm)+'^PQ1,0,1,N^XZ')

def make_zpl(items,dpi=203,gap_mm=GAP_MM,pitch_mm=PITCH_MM):
    labels=list(expanded(items))
    return ''.join(zpl_row(labels[i],labels[i+1] if i+1<len(labels) else None,dpi,gap_mm,pitch_mm)
                   for i in range(0,len(labels),2))

def network_print(ip,port,data):
    with socket.create_connection((ip,int(port)),timeout=7) as sock:
        sock.settimeout(15)
        sock.sendall(data.encode('utf-8'))


def zpl_safe(text):
    # UTF-8 for Zebra printers with ^CI28 support
    return str(text).replace('^',' ').replace('~',' ').replace('\r',' ').replace('\n',' ')

def zpl_label(p,quantity,dpi):
    factor=dpi/25.4
    pos=lambda mm_value:round(mm_value*factor)
    width,height=pos(LABEL_W),pos(LABEL_H)
    desc=zpl_safe(p['description'])
    sku=zpl_safe(p['sku'])
    size=zpl_safe(p['size'] or '-')
    color=zpl_safe(p['color'] or '-')
    bc=p['barcode']
    # Use Code128 for reliable scanning of EAN/other numeric strings.
    # 45mm width: 13-digit Code128 fits with module width 2 at 203dpi.
    module=2 if dpi<=203 else 3
    barcode=f'^FO{pos(3)},{pos(7)}^BY{module},2,{pos(10)}^BCN,{pos(10)},Y,N,N^FD{zpl_safe(bc)}^FS'
    return (f'^XA^CI28^PW{width}^LL{height}^LH0,0^LS0'
        f'^FO{pos(1)},{pos(1)}^A0N,{pos(2.7)},{pos(2.0)}^FB{width-pos(2)},1,0,C^FD{desc}^FS'
        +barcode+
        f'^FO{pos(1)},{pos(22)}^A0N,{pos(2.3)},{pos(1.7)}^FB{width-pos(2)},1,0,C^FDSKU: {sku}^FS'
        f'^FO{pos(1)},{pos(26)}^A0N,{pos(2.4)},{pos(1.6)}^FDSize: {size}^FS'
        f'^FO{pos(17)},{pos(26)}^A0N,{pos(2.3)},{pos(1.5)}^FDColor: {color}^FS'
        f'^PQ{quantity},0,1,N^XZ')

def list_zebra_printers():
    try:
        import win32print
        flags=win32print.PRINTER_ENUM_LOCAL|win32print.PRINTER_ENUM_CONNECTIONS
        return [entry[2] for entry in win32print.EnumPrinters(flags) if 'zebra' in entry[2].lower() or 'zt230' in entry[2].lower()]
    except ImportError:return []

def print_zpl(printer,items,dpi):
    try:import win32print
    except ImportError:raise RuntimeError('Voor rechtstreeks printen ontbreekt pywin32. Voer START.bat opnieuw uit.')
    if not printer:raise ValueError('Selecteer eerst de Zebra-printer.')
    handle=win32print.OpenPrinter(printer)
    try:
        job=win32print.StartDocPrinter(handle,1,('The Vandal stickers',None,'RAW'))
        try:
            win32print.StartPagePrinter(handle)
            try:
                data=''.join(zpl_label(p,q,dpi) for p,q in items).encode('utf-8')
                win32print.WritePrinter(handle,data)
            finally:win32print.EndPagePrinter(handle)
        finally:win32print.EndDocPrinter(handle)
    finally:win32print.ClosePrinter(handle)

class App:
    def __init__(self,root):
        self.root=root;root.title('The Vandal — ZT230 stickerprogramma Mac v7 (2 naast elkaar)');root.geometry('1190x790');root.minsize(990,630)
        self.products=[];self.visible=[];self.cart={};self.last_pdf=None
        saved=load_settings()
        st=ttk.Style();st.theme_use('clam')
        st.configure('Treeview',rowheight=27)
        outer=ttk.Frame(root,padding=14);outer.pack(fill='both',expand=True)
        top=ttk.Frame(outer);top.pack(fill='x')
        ttk.Label(top,text='THE VANDAL  |  STICKERPRINTER',font=('Arial',17,'bold')).pack(side='left')
        ttk.Button(top,text='Excel-bestand openen',command=self.choose_excel).pack(side='right')
        self.file_label=ttk.Label(outer,text='');self.file_label.pack(anchor='w',pady=(8,12))
        ttk.Label(outer,text='Zoeken op omschrijving, SKU, barcode, maat of kleur:').pack(anchor='w')
        self.query=tk.StringVar();entry=ttk.Entry(outer,textvariable=self.query,font=('Arial',12));entry.pack(fill='x',pady=(4,10));self.query.trace_add('write',lambda *_:self.refresh())
        # Printerbalk BOVEN de productlijsten: blijft altijd zichtbaar, ook op kleinere schermen.
        printerbar=ttk.LabelFrame(outer,text='Printerinstellingen — Zebra ZT230 (203 dpi, 2 × 45 × 30 mm)',padding=8)
        printerbar.pack(fill='x',pady=(0,8))
        settings=ttk.Frame(printerbar);settings.pack(fill='x')
        ttk.Label(settings,text='IP-adres:').pack(side='left')
        self.ip=tk.StringVar(value=saved['ip'])
        ttk.Entry(settings,textvariable=self.ip,width=16).pack(side='left',padx=5)
        ttk.Label(settings,text='Poort:').pack(side='left')
        self.port=tk.StringVar(value=saved['port'])
        ttk.Entry(settings,textvariable=self.port,width=6).pack(side='left',padx=5)
        ttk.Label(settings,text='Tussenruimte (mm):').pack(side='left',padx=(9,3))
        self.gap=tk.StringVar(value=saved['gap'])
        ttk.Entry(settings,textvariable=self.gap,width=5).pack(side='left')
        ttk.Label(settings,text='Rijafstand (mm):').pack(side='left',padx=(9,3))
        self.pitch=tk.StringVar(value=saved['pitch'])
        ttk.Entry(settings,textvariable=self.pitch,width=5).pack(side='left')
        actions=ttk.Frame(printerbar);actions.pack(fill='x',pady=(8,0))
        ttk.Button(actions,text='TESTAFDRUK (2 etiketten)',command=self.test_print).pack(side='left',padx=(0,10))
        ttk.Button(actions,text='PRINT SELECTIE NAAR ZEBRA',command=self.direct_print).pack(side='left',padx=(0,10))
        ttk.Button(actions,text='Maak PDF',command=self.export).pack(side='left')
        ttk.Button(actions,text='Printerinstellingen opslaan',command=self.save_settings).pack(side='left',padx=(10,0))
        paned=ttk.Panedwindow(outer,orient='vertical');paned.pack(fill='both',expand=True)
        a=ttk.LabelFrame(paned,text='Producten met barcode (dubbelklik om toe te voegen)',padding=7);paned.add(a,weight=3)
        self.products_tree=self.tree(a,('sku','description','size','color','barcode'),('SKU','Omschrijving','Maat','Kleur','Barcode'),(155,345,65,130,160))
        controls=ttk.Frame(a);controls.pack(fill='x',pady=(8,0))
        ttk.Label(controls,text='Aantal:').pack(side='left')
        self.qty=tk.StringVar(value='1');ttk.Entry(controls,textvariable=self.qty,width=9,font=('Arial',12)).pack(side='left',padx=7)
        ttk.Button(controls,text='Voeg toe met ingevuld aantal',command=self.add).pack(side='left')
        ttk.Label(controls,text='Typ een aantal (1–9999) en druk Enter').pack(side='left',padx=12)
        entry.bind('<Return>',lambda e:self.add())
        self.count_label=ttk.Label(controls,text='');self.count_label.pack(side='right')
        b=ttk.LabelFrame(paned,text='Afdruklijst',padding=7);paned.add(b,weight=2)
        self.cart_tree=self.tree(b,('sku','description','quantity','barcode'),('SKU','Omschrijving','Aantal','Barcode'),(165,420,75,180))
        buttons=ttk.Frame(b);buttons.pack(fill='x',pady=(8,0))
        ttk.Button(buttons,text='Aantal aanpassen (invullen)',command=self.edit_qty).pack(side='left',padx=(0,5))
        ttk.Button(buttons,text='Verwijder selectie',command=self.remove).pack(side='left',padx=5)
        ttk.Button(buttons,text='LEEG HELE SELECTIE',command=self.clear).pack(side='left',padx=5)
        self.total_label=ttk.Label(buttons,text='0 stickers');self.total_label.pack(side='right')
        self.products_tree.bind('<Double-1>',lambda e:self.add())
        self.cart_tree.bind('<Double-1>',self.edit_cart_cell)
        previous=get_last_excel()
        if previous:self.open_excel(previous)
        else:self.file_label.configure(text='Geen Excel-bestand geladen. Klik op Excel-bestand openen.')
    def tree(self,parent,cols,labels,widths):
        wrap=ttk.Frame(parent);wrap.pack(fill='both',expand=True)
        tree=ttk.Treeview(wrap,columns=cols,show='headings',selectmode='browse')
        for col,label,width in zip(cols,labels,widths):tree.heading(col,text=label);tree.column(col,width=width,minwidth=55,stretch=True)
        scroll=ttk.Scrollbar(wrap,orient='vertical',command=tree.yview);tree.configure(yscrollcommand=scroll.set)
        tree.pack(side='left',fill='both',expand=True);scroll.pack(side='right',fill='y');return tree
    def choose_excel(self):
        path=filedialog.askopenfilename(filetypes=[('Excel-bestanden','*.xlsx'),('Alle bestanden','*.*')])
        if path:self.open_excel(path)
    def open_excel(self,path):
        try:
            self.products=load_products(path);self.cart={};self.file_label.configure(text=f'Bestand: {path}  |  {len(self.products)} producten');self.refresh();self.refresh_cart();remember_excel(path)
        except Exception as exc:messagebox.showerror('Excel-fout',str(exc))
    def refresh(self):
        term=self.query.get().casefold().strip()
        self.visible=[(i,p) for i,p in enumerate(self.products) if p['barcode'].strip() and all(t in ' '.join(p.values()).casefold() for t in term.split())][:500]
        self.products_tree.delete(*self.products_tree.get_children())
        for i,p in self.visible:self.products_tree.insert('', 'end',iid=str(i),values=(p['sku'],p['description'],p['size'],p['color'],p['barcode'] or '—'))
        self.count_label.configure(text=f'{len(self.visible)} getoond / {sum(bool(p["barcode"].strip()) for p in self.products)} met barcode (max. 500)')
    def selected_qty(self):
        try:
            q=int(self.qty.get())
            if q<1 or q>9999:raise ValueError
            return q
        except ValueError:messagebox.showwarning('Aantal','Vul een geheel aantal tussen 1 en 9999 in.');return None
    def add(self):
        sel=self.products_tree.selection()
        if not sel:return
        i=int(sel[0]);p=self.products[i]
        if not p['barcode']:
            messagebox.showwarning('Geen barcode','Dit product heeft geen barcode in Excel en kan niet worden toegevoegd.');return
        q=self.selected_qty()
        if q is None:return
        self.cart[i]=self.cart.get(i,0)+q;self.refresh_cart()
    def refresh_cart(self):
        self.cart_tree.delete(*self.cart_tree.get_children())
        for i,q in self.cart.items():
            p=self.products[i];self.cart_tree.insert('','end',iid=str(i),values=(p['sku'],p['description'],q,p['barcode']))
        self.total_label.configure(text=f'{sum(self.cart.values())} stickers / {len(self.cart)} producten')
    def edit_cart_cell(self,event):
        # Dubbelklik op de kolom Aantal: wijzig het aantal direct in de afdruklijst.
        row=self.cart_tree.identify_row(event.y)
        col=self.cart_tree.identify_column(event.x)
        if not row:return
        self.cart_tree.selection_set(row)
        if col!='#3':
            self.edit_qty()
            return
        bbox=self.cart_tree.bbox(row,'quantity')
        if not bbox:return
        x,y,w,h=bbox
        editor=ttk.Entry(self.cart_tree,font=('Arial',11))
        editor.insert(0,str(self.cart[int(row)]))
        editor.place(x=x,y=y,width=w,height=h)
        editor.focus_set();editor.select_range(0,'end')
        def finish(save=True):
            if not editor.winfo_exists():return
            raw=editor.get().strip()
            if save:
                try:
                    q=int(raw)
                    if not 1<=q<=9999:raise ValueError
                except ValueError:
                    messagebox.showwarning('Aantal','Vul een geheel aantal van 1 tot 9999 in.')
                    editor.focus_set();return
                self.cart[int(row)]=q
            editor.destroy();self.refresh_cart()
        editor.bind('<Return>',lambda e:finish())
        editor.bind('<Escape>',lambda e:finish(False))
        editor.bind('<FocusOut>',lambda e:finish())

    def edit_qty(self):
        sel=self.cart_tree.selection()
        if not sel:return
        current=self.cart[int(sel[0])]
        answer=tk.simpledialog.askinteger('Aantal stickers','Vul het gewenste aantal stickers in:',initialvalue=current,minvalue=1,maxvalue=9999,parent=self.root)
        if answer is not None:self.cart[int(sel[0])]=answer;self.refresh_cart()
    def remove(self):
        for i in self.cart_tree.selection():self.cart.pop(int(i),None)
        self.refresh_cart()
    def clear(self):
        if not self.cart:return
        if messagebox.askyesno('Selectie leegmaken','Alle producten uit de afdruklijst verwijderen?'):
            self.cart.clear();self.refresh_cart()
    def save_settings(self,show_message=True):
        try:
            ip,port,gap,pitch=self.print_settings()
            SETTINGS_DIR.mkdir(parents=True,exist_ok=True)
            settings={'ip':ip,'port':str(port),'gap':str(gap),'pitch':str(pitch)}
            temp=SETTINGS_FILE.with_suffix('.tmp')
            temp.write_text(json.dumps(settings,indent=2),encoding='utf-8')
            temp.replace(SETTINGS_FILE)
            if show_message:messagebox.showinfo('Printerinstellingen','Instellingen opgeslagen voor de volgende keer.\n'+str(SETTINGS_FILE))
        except Exception as exc:
            if show_message:messagebox.showerror('Instellingen niet opgeslagen',str(exc))
    
    def print_settings(self):
        ip=self.ip.get().strip()
        if not ip:raise ValueError('Vul het IP-adres van de printer in.')
        ipaddress.ip_address(ip)
        port=int(self.port.get())
        gap=float(self.gap.get().replace(',','.'))
        pitch=float(self.pitch.get().replace(',','.'))
        if not (1<=port<=65535 and 0<=gap<=10 and 30<=pitch<=45):
            raise ValueError('Controleer poort, tussenruimte en rijafstand.')
        return ip,port,gap,pitch
    def test_print(self):
        if not messagebox.askyesno('Testafdruk','Twee testetiketten naast elkaar afdrukken?\nEr wordt 1 rij verbruikt.'):
            return
        try:
            ip,port,gap,pitch=self.print_settings()
            a={'description':'TEST LINKS','sku':'TEST-L','size':'L','color':'Blue','barcode':'6151103755701'}
            b={'description':'TEST RECHTS','sku':'TEST-R','size':'M','color':'Black','barcode':'6151103755701'}
            network_print(ip,port,zpl_row(a,b,203,gap,pitch))
            self.save_settings(show_message=False)
            messagebox.showinfo('Verzonden','Testafdruk verzonden. Controleer uitlijning en scanbaarheid.')
        except Exception as exc:messagebox.showerror('Printerfout',str(exc))
    def direct_print(self):
        if not self.cart:messagebox.showinfo('Afdruklijst','Voeg eerst producten toe.');return
        total=sum(self.cart.values())
        if not messagebox.askyesno('Bevestig afdruk',f'{total} stickers / {(total+1)//2} rijen afdrukken via het netwerk?'):return
        try:
            ip,port,gap,pitch=self.print_settings()
            data=make_zpl([(self.products[i],q) for i,q in self.cart.items()],203,gap,pitch)
            network_print(ip,port,data)
            self.save_settings(show_message=False)
            messagebox.showinfo('Verzonden',f'{total} stickers naar {ip}:{port} verzonden. Controleer de printer.')
        except Exception as exc:messagebox.showerror('Printerfout',str(exc))
    def export(self):
        if not self.cart:messagebox.showinfo('Afdruklijst','Voeg eerst producten toe.');return
        path=filedialog.asksaveasfilename(defaultextension='.pdf',initialfile='The_Vandal_stickers.pdf',filetypes=[('PDF','*.pdf')])
        if not path:return
        try:
            make_pdf([(self.products[i],q) for i,q in self.cart.items()],path)
            if messagebox.askyesno('PDF aangemaakt',f'{sum(self.cart.values())} stickers opgeslagen in:\n{path}\n\nPDF nu openen om af te drukken?'):
                if sys.platform=='win32':os.startfile(path)
                elif sys.platform=='darwin':subprocess.Popen(['open',path])
                else:subprocess.Popen(['xdg-open',path])
        except Exception as exc:messagebox.showerror('Fout bij aanmaken PDF',str(exc))

if __name__=='__main__':
    root=tk.Tk();App(root);root.mainloop()
